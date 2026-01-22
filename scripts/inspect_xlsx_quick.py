#!/usr/bin/env python3

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter


MONTH_RE = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\b", re.I)


@dataclass(frozen=True)
class Hit:
    sheet: str
    cell: str
    value: str


def _norm_cell(value: object | None) -> str:
    if value is None:
        return ""
    # Reduce float noise for integer-like cells.
    if isinstance(value, float) and abs(value - round(value)) < 1e-9:
        return str(int(round(value)))
    return str(value).strip()


def _scan_sheet(
    sheet_name: str,
    ws,
    *,
    max_rows: int,
    max_cols: int,
    needles: list[re.Pattern[str]],
) -> tuple[list[dict], list[Hit]]:
    month_header_candidates: list[dict] = []
    hits: list[Hit] = []

    for r in range(1, min(max_rows, ws.max_row or 1) + 1):
        row_vals: list[str] = []

        for c in range(1, min(max_cols, ws.max_column or 1) + 1):
            s = _norm_cell(ws.cell(row=r, column=c).value)
            row_vals.append(s)

            if s and any(rx.search(s) for rx in needles):
                hits.append(
                    Hit(
                        sheet=sheet_name,
                        cell=f"{get_column_letter(c)}{r}",
                        value=(s[:160] + "…") if len(s) > 160 else s,
                    )
                )

        month_cols = [idx for idx, s in enumerate(row_vals, start=1) if s and MONTH_RE.match(s)]
        if len(month_cols) >= 3:
            month_header_candidates.append({"row": r, "month_cols": month_cols})

    return month_header_candidates, hits


def _chunk(iterable: Iterable[Hit], limit: int) -> list[Hit]:
    out: list[Hit] = []
    for item in iterable:
        out.append(item)
        if len(out) >= limit:
            break
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description="Quick XLSX keyword + month-header inspection")
    parser.add_argument("path", help="Path to .xlsx file")
    parser.add_argument("--rows", type=int, default=60, help="Rows to scan per sheet")
    parser.add_argument("--cols", type=int, default=30, help="Cols to scan per sheet")
    parser.add_argument(
        "--needles",
        default="revenue,sales,cogs,expense,opex,cash,runway,burn,ebitda,profit,valuation",
        help="Comma-separated keywords/regex fragments to search for",
    )
    parser.add_argument("--max-hits", type=int, default=30, help="Max keyword hits to print per sheet")
    parser.add_argument(
        "--sheet",
        default="",
        help="Optional: only inspect a specific sheet name (exact match)",
    )
    parser.add_argument(
        "--grid-rows",
        type=int,
        default=0,
        help="If >0, print a simple preview grid for the first N rows",
    )
    parser.add_argument(
        "--grid-cols",
        type=int,
        default=0,
        help="If >0, print a simple preview grid for the first N cols",
    )

    args = parser.parse_args()

    needles = [re.compile(fragment.strip(), re.I) for fragment in args.needles.split(",") if fragment.strip()]

    wb = load_workbook(args.path, data_only=True)
    print({"path": args.path, "sheets": wb.sheetnames})

    sheet_names = wb.sheetnames
    if args.sheet:
        sheet_names = [name for name in sheet_names if name == args.sheet]
        if not sheet_names:
            available = ", ".join(wb.sheetnames)
            raise SystemExit(f"Sheet not found: {args.sheet}. Available: {available}")

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        month_headers, hits = _scan_sheet(
            sheet_name,
            ws,
            max_rows=args.rows,
            max_cols=args.cols,
            needles=needles,
        )
        print(f"\n=== SHEET: {sheet_name} ===")
        print({
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "month_header_candidates": month_headers[:5],
            "keyword_hits": [h.__dict__ for h in _chunk(hits, args.max_hits)],
        })

        if args.grid_rows > 0 and args.grid_cols > 0:
            grid_rows = min(args.grid_rows, ws.max_row or 1)
            grid_cols = min(args.grid_cols, ws.max_column or 1)
            print(f"-- GRID (first {grid_rows} rows x {grid_cols} cols) --")
            for r in range(1, grid_rows + 1):
                row_vals = [_norm_cell(ws.cell(row=r, column=c).value) for c in range(1, grid_cols + 1)]
                # Keep the preview compact.
                row_vals = [(v[:32] + "…") if len(v) > 32 else v for v in row_vals]
                print("\t".join(row_vals))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
